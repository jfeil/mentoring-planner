"""Data model and Excel reading for the mentoring planner.

The workbook is the input format. Three sheets are read:

* groups   -- one row per class-split group: name + which class it comes from
* peers    -- one row per mentoring peer: name + topic + the classes its pupils are in
* calendar -- class x week grid marking the weeks each class is on a trip

Everything else (open weeks, feasible group/peer pairs) is derived.
"""

from __future__ import annotations

import datetime
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook as _open_workbook

GROUP_SHEETS = ("groups", "group", "gruppen", "gruppe")
PEER_SHEETS = ("peers", "peer", "betreuer", "mentor", "mentoren", "mentees")
CALENDAR_SHEETS = ("calendar", "kalender", "trips", "trips_calendar", "ausfluege", "termine")
SETTINGS_SHEETS = ("settings", "optionen", "einstellungen")

GROUP_HEADERS = {
    "group": ("group", "groups", "gruppe", "gruppen"),
    "cls": ("class", "classes", "klasse", "klassen"),
}
PEER_HEADERS = {
    "name": ("peers", "peer", "peers surnames", "surnames", "names", "name", "members",
             "betreuer", "mentoren", "mentor", "teilnehmer", "nachnamen"),
    "topic": ("topic", "topics", "thema", "themenschwerpunkt", "schwerpunkt"),
    "classes": ("classes", "class", "klassen", "klasse"),
    "weeks": (
        "available weeks",
        "weeks available",
        "usable weeks",
        "weeks",
        "week",
        "wochen",
        "woche",
        "verfgbare wochen",
        "kw",
    ),
    "max_sessions": ("max sessions", "maxsessions", "max", "sessions max", "einsaetze max",
                     "max einsatze", "max einsaetze", "limit"),
}
SETTINGS_KEYS = {
    "blackouts": ("blackout weeks", "blackouts", "blocked weeks", "blockierte wochen", "ferien"),
    "peer_cap": ("max sessions per peer", "max einsatze pro peer", "sitzungslimit"),
    "group_week_cap": (
        "max sessions per group per week",
        "sessions per group per week",
        "max pro gruppe und woche",
        "sitzungen pro gruppe und woche",
    ),
}
TOPIC_ORDER_SHEETS = ("topics", "themen")

# A calendar header cell: 36, W36 or KW36. Anchored, so class names like '6a' are not weeks.
WEEK_HEADER = re.compile(r"(?:kw|w)?\s*(\d{1,2})")

TRUE_WORDS = {"x", "a", "j", "ja", "y", "yes", "true", "1", "trip", "out", "untericht", "ausflug", "fehlt"}


class WorkbookError(Exception):
    """The workbook cannot be understood; message is meant for the end user."""


@dataclass(eq=False)
class Group:
    """A class-split group, e.g. A1 out of class 5a. Compared by identity."""

    name: str
    cls: str
    assigned: list["Assignment"] = field(default_factory=list, repr=False)


@dataclass(eq=False)
class Peer:
    """One mentoring peer: a fixed group of older pupils serving one topic.

    ``members`` are the pupils' surnames as written in the sheet; ``name`` is the
    label used in the reports, e.g. "Sander; Weise".
    """

    name: str
    topic: str
    classes: tuple[str, ...] = ()
    members: tuple[str, ...] = ()
    weeks: frozenset[int] | None = None  # None = available every week
    max_sessions: int | None = None  # None = workbook default, or no limit at all
    assigned: list["Assignment"] = field(default_factory=list, repr=False)

    def trip_weeks(self, trips: dict[str, set[int]]) -> set[int]:
        out: set[int] = set()
        for cls in self.classes:
            out |= trips.get(cls, set())
        return out


@dataclass
class Assignment:
    group: Group
    peer: Peer
    week: int


@dataclass
class Data:
    """Everything the planner needs, plus the derived feasibility sets."""

    groups: list[Group]
    peers: list[Peer]
    classes: tuple[str, ...]
    weeks: tuple[int, ...]
    trips: dict[str, set[int]]
    topics: tuple[str, ...]
    blackouts: frozenset[int] = frozenset()  # holiday / exam weeks: never any mentoring
    peer_cap: int | None = None  # default 'Max sessions' for peers that leave it blank
    group_week_cap: int | None = None  # sessions one group may have in one week; None = no limit

    @property
    def open_weeks(self) -> tuple[int, ...]:
        """Weeks with a room: a tutored class has to be away."""

        return tuple(w for w in self.weeks if self.is_open(w))

    @property
    def tutored_classes(self) -> tuple[str, ...]:
        """The classes the groups are split from; peer classes are not part of this."""

        seen: list[str] = []
        for group in self.groups:
            if group.cls not in seen:
                seen.append(group.cls)
        return tuple(seen)

    def is_open(self, week: int) -> bool:
        """The room exists only in weeks where a tutored class is on a trip.

        A peer's class being away does not open a room; it only makes that peer
        unavailable.
        """

        if week in self.blackouts:
            return False
        return any(week in self.trips.get(cls, set()) for cls in self.tutored_classes)

    def cap_of(self, peer: Peer) -> int | None:
        """How many sessions this peer may take; None means no limit."""

        return peer.max_sessions if peer.max_sessions is not None else self.peer_cap

    def allowed_weeks(self, group: Group, peer: Peer) -> tuple[int, ...]:
        """Weeks where this group could meet this peer."""

        blocked = self.trips.get(group.cls, set()) | peer.trip_weeks(self.trips)
        weeks = [w for w in self.weeks if w not in blocked and self.is_open(w)]
        if peer.weeks is not None:
            weeks = [w for w in weeks if w in peer.weeks]
        return tuple(weeks)

    def peers_for(self, topic: str) -> list[Peer]:
        return [p for p in self.peers if p.topic == topic]

    def weeks_left_for(self, group: Group, peer: Peer, used: dict[int, int]) -> tuple[int, ...]:
        """Weeks this group could still use for this peer, given its current week loads."""

        cap = self.group_week_cap
        return tuple(
            w
            for w in self.allowed_weeks(group, peer)
            if cap is None or used.get(w, 0) < cap
        )


# ---------------------------------------------------------------- loading ----


def load_workbook(path: Path) -> Data:
    wb = _open_workbook(path, data_only=True)
    problems: list[str] = []

    groups_sheet = _find_sheet(wb, GROUP_SHEETS, problems, "a groups sheet (Group | Class)")
    peers_sheet = _find_sheet(wb, PEER_SHEETS, problems, "a peers sheet (Peer | Topic | Classes)")
    calendar_sheet = _find_sheet(wb, CALENDAR_SHEETS, problems, "a calendar sheet (class x week)")
    if problems:
        raise WorkbookError("\n".join(problems))

    classes, trips, week_numbers = _read_calendar(wb[calendar_sheet], problems)
    groups = _read_groups(wb[groups_sheet], classes, problems)
    peers, topics = _read_peers(wb[peers_sheet], classes, problems)
    blackouts, peer_cap, group_week_cap = _read_settings(wb, problems)

    weeks = tuple(sorted(set(week_numbers) | {w for s in trips.values() for w in s}))
    tutored = {g.cls for g in groups}
    if groups and weeks and not any(trips.get(cls) for cls in tutored):
        problems.append(
            "none of the classes the groups come from is ever on a trip, so a room is never "
            "free and no session can be planned - mark the trips in the calendar sheet"
        )
    elif len(tutored) == 1 and groups:
        (only,) = sorted(tutored)
        problems.append(
            f"all groups come from the one class '{only}', and a room only opens while a "
            "tutored class travels - so that class is away in every week a room is free and "
            "none of its groups can ever meet; list a second class that travels too"
        )

    if problems:
        raise WorkbookError("\n".join(problems))
    return Data(
        groups=groups,
        peers=peers,
        classes=classes,
        weeks=weeks,
        trips=trips,
        topics=topics,
        blackouts=frozenset(blackouts & set(weeks)),
        peer_cap=peer_cap,
        group_week_cap=group_week_cap,
    )


def _read_settings(wb, problems: list[str]) -> tuple[set[int], int | None, int | None]:
    """Optional sheet of 'Setting | Value' lines: blackout weeks and the session limits."""

    available = {_norm(name): name for name in wb.sheetnames}
    name = next((available[key] for key in SETTINGS_SHEETS if key in available), None)
    if name is None:
        return set(), None, None
    ws = wb[name]
    found: dict[str, str] = {}
    for row in range(1, ws.max_row + 1):
        key = _norm(_cell(ws, row, 1))
        value = _cell(ws, row, 2)
        if not key or not value:
            continue
        for setting, aliases in SETTINGS_KEYS.items():
            if key in [_norm(alias) for alias in aliases]:
                found[setting] = value
    blackouts: set[int] = set()
    peer_cap: int | None = None
    group_week_cap: int | None = None
    if "blackouts" in found:
        blackouts = set(_read_weeks(found["blackouts"]))
    if "peer_cap" in found:
        peer_cap = _positive_int(found["peer_cap"], "Max sessions per peer", name, problems)
    if "group_week_cap" in found:
        group_week_cap = _positive_int(
            found["group_week_cap"], "Max sessions per group per week", name, problems
        )
    return blackouts, peer_cap, group_week_cap


def _positive_int(text: str, label: str, sheet: str, problems: list[str]) -> int | None:
    if not text.isdigit() or int(text) == 0:
        problems.append(f"'{sheet}': '{label}' needs a whole number above 0, not '{text}'")
        return None
    return int(text)


def _find_sheet(wb, aliases: Iterable[str], problems: list[str], what: str) -> str | None:
    available = {_norm(name): name for name in wb.sheetnames}
    for alias in aliases:
        if alias in available:
            return available[alias]
    for name in wb.sheetnames:
        if any(alias in name for alias in aliases):
            return wb[name].title
    problems.append(f"the workbook needs {what}; it has {', '.join(wb.sheetnames)}")
    return None


def _norm(text: str) -> str:
    text = text.strip().lower().replace("ue", "u").replace("ä", "a")
    text = text.replace("ö", "o").replace("ü", "u").replace("ß", "s")
    return re.sub(r"[^a-z0-9]+", "", text)


def _header_row(ws, required: Iterable[str], optional: Iterable[str], sheet: str):
    """Locate the header line and return it as {key: column index}."""

    wanted = {key: [_norm(a) for a in aliases] for key, aliases in required.items()}
    wanted.update({key: [_norm(a) for a in aliases] for key, aliases in optional.items()})
    for row in range(1, min(ws.max_row, 8) + 1):
        values = {c: _norm(str(ws.cell(row, c).value)) for c in range(1, ws.max_column + 1)}
        found: dict[str, int] = {}
        for col, value in values.items():
            for key, aliases in wanted.items():
                if value in aliases and key not in found:
                    found[key] = col
        if all(key in found for key in required):
            return row, found
    labels = ", ".join(next(iter(a)) for a in required.values())
    raise WorkbookError(f"the '{sheet}' sheet needs header columns for: {labels}")


def _read_groups(ws, classes: tuple[str, ...], problems: list[str]) -> list[Group]:
    header, cols = _header_row(ws, GROUP_HEADERS, {}, "groups")
    groups: list[Group] = []
    seen: set[str] = set()
    for row in range(header + 1, ws.max_row + 1):
        name = _cell(ws, row, cols["group"])
        cls = _cell(ws, row, cols["cls"])
        if not name and not cls:
            continue
        if not name or not cls:
            problems.append(f"groups row {row}: needs both a group name and a class")
            continue
        if name in seen:
            problems.append(f"groups row {row}: group '{name}' is listed twice")
            continue
        if cls not in classes:
            problems.append(
                f"groups row {row}: group '{name}' comes from class '{cls}', which the "
                "calendar does not list - a group's class is what opens and blocks weeks"
            )
        seen.add(name)
        groups.append(Group(name=name, cls=cls))
    if not groups:
        problems.append("the groups sheet lists no groups")
    return groups


def _read_peers(ws, classes: tuple[str, ...], problems: list[str]):
    required = {key: aliases for key, aliases in PEER_HEADERS.items() if key not in ("weeks", "max_sessions")}
    optional = {key: PEER_HEADERS[key] for key in ("weeks", "max_sessions")}
    header, cols = _header_row(ws, required, optional, "peers")
    peers: list[Peer] = []
    seen: set[str] = set()
    topics: list[str] = []
    for row in range(header + 1, ws.max_row + 1):
        raw_members = _cell(ws, row, cols["name"])
        topic = _cell(ws, row, cols["topic"])
        raw_classes = _cell(ws, row, cols["classes"])
        if not raw_members and not topic and not raw_classes:
            continue
        if not raw_members or not topic or not raw_classes:
            problems.append(
                f"peers row {row}: needs the pupils' surnames, a topic and their classes"
            )
            continue
        members = _members(raw_members)
        name = "; ".join(members)
        if name in seen:
            problems.append(f"peers row {row}: peer '{name}' is listed twice")
            continue
        seen.add(name)
        peer_classes = _split(raw_classes)
        unknown = [c for c in peer_classes if c not in classes]
        if unknown:
            problems.append(
                f"peers row {row}: peer '{name}' mentions class(es) "
                f"{', '.join(sorted(set(unknown)))} that the calendar does not know"
            )
        raw_weeks = _cell(ws, row, cols["weeks"]) if "weeks" in cols else ""
        weeks = _read_weeks(raw_weeks) if raw_weeks else None  # blank = every week
        limit = _cell(ws, row, cols["max_sessions"]) if "max_sessions" in cols else ""
        if limit and not limit.isdigit():
            problems.append(f"peers row {row}: 'Max sessions' for '{name}' must be a number")
            continue
        if topic not in topics:
            topics.append(topic)
        peers.append(
            Peer(
                name=name,
                topic=topic,
                classes=peer_classes,
                members=members,
                weeks=weeks,
                max_sessions=int(limit) if limit else None,
            )
        )
    if not peers:
        problems.append("the peers sheet lists no peers")
    return peers, tuple(topics)


def _read_calendar(ws, problems: list[str]):
    """Class x week grid: rows are classes, columns are weeks."""

    week_numbers: list[int] = []
    header_row = 1
    for row in range(1, min(ws.max_row, 8) + 1):
        found = [
            (col, _week_header(ws.cell(row, col).value))
            for col in range(2, ws.max_column + 1)
        ]
        marks = sum(
            1
            for col in range(2, ws.max_column + 1)
            if _norm(str(ws.cell(row, col).value or "")) in TRUE_WORDS
        )
        weeks_found = sum(v is not None for _, v in found)
        if weeks_found >= 1 and marks == 0 and _week_header(ws.cell(row, 1).value) is None:
            header_row = row
            week_numbers = [v for _, v in found if v is not None]
            break
    if not week_numbers:
        raise WorkbookError(
            f"the '{ws.title}' sheet needs a header row of week numbers starting in column 2 "
            "(for example 36, 37, 38 ...)"
        )

    classes: list[str] = []
    trips: dict[str, set[int]] = {}
    for row in range(header_row + 1, ws.max_row + 1):
        cls = _cell(ws, row, 1)
        if not cls:
            continue
        if cls in trips:
            problems.append(f"'{ws.title}' row {row}: class '{cls}' is listed twice")
            continue
        classes.append(cls)
        on_trip: set[int] = set()
        for col, week in zip(range(2, ws.max_column + 1), week_numbers):
            value = ws.cell(row, col).value
            if value is None:
                continue
            text = _norm(str(value))
            if text in TRUE_WORDS or (week is not None and str(value).strip() in {str(week)}):
                on_trip.add(week)
        trips[cls] = on_trip
    if not classes:
        problems.append(f"the '{ws.title}' sheet lists no classes")
    return tuple(classes), trips, week_numbers


def _week_header(value) -> int | None:
    """A calendar header cell: 36, W36, KW36 or a date; anything else is not a week."""

    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isocalendar()[1]
    if isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
        return int(value)
    match = WEEK_HEADER.fullmatch(_norm(str(value)))
    return int(match.group(1)) if match else None


def _cell(ws, row: int, col: int) -> str:
    value = ws.cell(row, col).value
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split(text: str) -> tuple[str, ...]:
    parts = re.split(r"[,/;+&\s]+", text)
    out: list[str] = []
    for part in parts:
        part = part.strip()
        if part and part not in out:
            out.append(part)
    return tuple(out)


def _members(text: str) -> tuple[str, ...]:
    """The pupils of a peer, written 'Sander; Weise' (inner spaces are kept)."""

    parts = re.split(r"[;/,\n]+", text)
    out: list[str] = []
    for part in parts:
        part = " ".join(part.split())
        if part and part not in out:
            out.append(part)
    return tuple(out)


def _read_weeks(text: str) -> frozenset[int]:
    weeks: set[int] = set()
    for part in _split(text):
        match = re.fullmatch(r"(\d{1,2})\s*[-.]\s*(\d{1,2})", part)
        if match:
            start, end = int(match.group(1)), int(match.group(2))
            weeks.update(range(min(start, end), max(start, end) + 1))
        elif part.isdigit():
            weeks.add(int(part))
    return frozenset(weeks)
