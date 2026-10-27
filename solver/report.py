"""Writes the plan back into the workbook.

Result sheets: Plan (the grid people read), Sessions (one row per session),
Loads (how much each peer got), Overview and Notes.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from solver.model import Data
from solver.solver import Plan

RESULT_SHEETS = ("Plan", "Sessions", "Loads", "Overview", "Notes")

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
WEEK_FILL = PatternFill("solid", fgColor="F2F2F2")
BLOCKED_FILL = PatternFill("solid", fgColor="FCE4E4")


def write_plan(workbook: Path, data: Data, plan: Plan, out: Path | None = None) -> Path:
    target = out or workbook
    wb = load_workbook(workbook)
    for name in RESULT_SHEETS:
        for existing in [s for s in wb.sheetnames if s.lower() == name.lower()]:
            del wb[existing]

    _sessions_sheet(wb, plan)
    _plan_grid(wb, data, plan)
    _loads_sheet(wb, data, plan)
    _overview_sheet(wb, data, plan)
    _notes_sheet(wb, data, plan)
    wb.active = wb["Plan"]
    wb.save(target)
    return target


def _sheet(wb, name: str):
    ws = wb.create_sheet(name)
    return ws


def _header(ws, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(1, col, label)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _sessions_sheet(wb, plan: Plan) -> None:
    ws = _sheet(wb, "Sessions")
    _header(
        ws,
        [
            "Week",
            "Group",
            "Class",
            "Topic",
            "Peer group",
            "Its classes",
            "Peer could also meet in",
            "Other peers for this topic",
        ],
    )
    for s in plan.sessions:
        options = plan.alternatives(s)
        others = " | ".join(
            f"{p.name} ({', '.join(f'W{x}' for x in w)})" for p, w in options
        )
        spare = plan.option_count(s) - len(options)
        if spare > 0:
            others += f"; +{spare} more"
        ws.append(
            [
                s.week,
                s.group.name,
                s.group.cls,
                s.peer.topic,
                s.peer.name,
                ", ".join(s.peer.classes),
                ", ".join(f"W{w}" for w in plan.week_options(s) if w != s.week),
                others,
            ]
        )
    _widths(ws, [8, 10, 8, 14, 18, 14, 22, 46])


def _plan_grid(wb, data: Data, plan: Plan) -> None:
    ws = _sheet(wb, "Plan")
    _header(ws, ["Group", "Class"] + [f"W{w}" for w in data.weeks] + ["Sessions", "Missing"])
    by_group: dict[str, dict[int, list[str]]] = {}
    for s in plan.sessions:
        by_group.setdefault(s.group.name, {}).setdefault(s.week, []).append(
            f"{s.peer.topic} / {s.peer.name}"
        )
    for group in data.groups:
        cells = by_group.get(group.name, {})
        row = [group.name, group.cls]
        styles = []
        for week in data.weeks:
            if week in cells:
                row.append("\n".join(cells[week]))
                styles.append(None)
            elif not data.is_open(week):
                row.append("blocked" if week in data.blackouts else "no room")
                styles.append(WEEK_FILL)
            elif week in data.trips.get(group.cls, set()):
                row.append("class away")
                styles.append(WEEK_FILL)
            else:
                row.append("")
                styles.append(None)
        missing = [topic for group_name, topic in plan.unresolved if group_name == group.name]
        row += [sum(len(entry) for entry in cells.values()), ", ".join(missing)]
        ws.append(row)
        r = ws.max_row
        for offset, fill in enumerate(styles):
            cell = ws.cell(r, 3 + offset)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fill is not None:
                cell.fill = fill
        if missing:
            ws.cell(r, len(row)).fill = BLOCKED_FILL
    _widths(ws, [10, 9] + [20] * len(data.weeks) + [9, 22])

    note = ws.cell(
        ws.max_row + 2,
        1,
        "One entry per group and topic; a week cell can hold several sessions. "
        "The day within the week is set by hand.",
    )
    note.font = Font(italic=True)


def _loads_sheet(wb, data: Data, plan: Plan) -> None:
    ws = _sheet(wb, "Loads")
    _header(ws, ["Peer", "Topic", "Sessions", "Limit", "Free", "Weeks used", "Groups"])
    loads = plan.peer_loads()
    for peer in sorted(data.peers, key=lambda p: (p.topic, -loads[p.name], p.name)):
        cap = data.cap_of(peer)
        weeks = ", ".join(f"W{s.week}" for s in sorted(peer.assigned, key=lambda s: s.week))
        ws.append(
            [
                peer.name,
                peer.topic,
                loads[peer.name],
                cap if cap is not None else "-",
                "-" if cap is None else cap - loads[peer.name],
                weeks,
                ", ".join(s.group.name for s in sorted(peer.assigned, key=lambda s: s.group.name)),
            ]
        )
    _widths(ws, [18, 14, 10, 8, 8, 22, 24])


def _overview_sheet(wb, data: Data, plan: Plan) -> None:
    ws = _sheet(wb, "Overview")
    _header(ws, ["Topic", "Peers", "Sessions", "Idlest", "Busiest", "Balance"])
    loads = plan.peer_loads()
    for topic in data.topics:
        peers = data.peers_for(topic)
        values = [loads[p.name] for p in peers]
        low, high = (min(values), max(values)) if values else (0, 0)
        ws.append(
            [
                topic,
                len(peers),
                sum(values),
                low,
                high,
                "even" if high - low <= 1 else f"spread of {high - low}",
            ]
        )
    ws.append([])
    ws.append(["Weeks with a room", ", ".join(f"W{w}" for w in data.open_weeks)])
    ws.append(["No mentoring", ", ".join(f"W{w}" for w in data.weeks if not data.is_open(w)) or "-"])
    ws.append(["Sessions per week", ", ".join(f"W{w}: {n}" for w, n in sorted(plan.week_loads().items()))])
    ws.append(["Total sessions", len(plan.sessions)])
    ws.append(["Tutored classes", ", ".join(data.tutored_classes)])
    default = "none" if data.peer_cap is None else str(data.peer_cap)
    individual = ", ".join(f"{p.name}: {p.max_sessions}" for p in data.peers if p.max_sessions is not None)
    ws.append(["Sessions per peer - default", default])
    ws.append(["Sessions per peer - individual", individual or "-"])
    ws.append(
        [
            "Sessions per group per week",
            "no limit" if data.group_week_cap is None else str(data.group_week_cap),
        ]
    )
    _widths(ws, [26, 10, 10, 8, 8, 18])


def _notes_sheet(wb, data: Data, plan: Plan) -> None:
    ws = _sheet(wb, "Notes")
    _header(ws, ["Note"])
    for note in plan.notes:
        ws.append([note])
    if plan.unresolved:
        ws.append([])
        ws.append(["Could not be planned"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for group, topic in plan.unresolved:
            ws.append([f"{group} / {topic}"])
    if not plan.complete:
        cell = ws.cell(ws.max_row + 2, 1, "These pairs have no open week in common; check the calendar.")
        cell.font = Font(italic=True)
    _widths(ws, [110])


def _widths(ws, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
