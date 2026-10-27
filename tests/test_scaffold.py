"""The scaffold generator: what it writes must be an ordinary input workbook."""

from __future__ import annotations

import datetime

import pytest
from openpyxl import load_workbook

from solver import load_workbook as load_data
from solver import solve
from solver.scaffold import (
    Spec,
    group_names,
    mondays,
    parse_classes,
    week_of,
    weeks_from_dates,
    write_scaffold,
)

SPEC = Spec(
    tutored=[("6a", 2), ("6b", 1)],
    peer_classes=["9a", "9b"],
    topics=["Reading", "Maths"],
    weeks=[36, 37, 38],
    peers_per_topic=2,
    # A template nobody marked would open no room at all, so the example trips stay.
    trips={"6a": [36, 37], "6b": [38]},
)


def template(tmp_path, spec=SPEC):
    return write_scaffold(tmp_path / "template.xlsx", spec)


def test_it_writes_the_input_sheets(tmp_path):
    names = load_workbook(template(tmp_path)).sheetnames
    assert names == ["Groups", "Peers", "Calendar", "Settings", "Read me"]


def test_the_groups_are_the_class_splits(tmp_path):
    rows = list(load_workbook(template(tmp_path))["Groups"].iter_rows(min_row=2, values_only=True))
    assert rows == [("6a.1", "6a"), ("6a.2", "6a"), ("6b.1", "6b")]
    assert group_names(SPEC) == rows


def test_the_calendar_lists_both_kinds_of_class(tmp_path):
    ws = load_workbook(template(tmp_path))["Calendar"]
    assert [row[0] for row in ws.iter_rows(min_row=2, values_only=True)] == ["6a", "6b", "9a", "9b"]
    assert [cell.value for cell in ws[1]] == ["Class", "W36", "W37", "W38"]


def test_the_blank_template_loads(tmp_path):
    """The generator may not produce a sheet the reader refuses."""

    data = load_data(template(tmp_path))
    assert [g.name for g in data.groups] == ["6a.1", "6a.2", "6b.1"]
    assert data.topics == ("Reading", "Maths")
    assert data.weeks == (36, 37, 38)
    assert all(p.name.startswith("<peer") for p in data.peers)


def test_trips_and_settings_survive_the_round_trip(tmp_path):
    spec = Spec(
        tutored=[("6a", 1), ("6b", 1)],
        peer_classes=["9a"],
        topics=["Reading"],
        weeks=[34, 35, 36],
        trips={"6a": [34], "6b": [35], "9a": [35]},
        blackouts=[36],
        peer_cap=3,
        group_week_cap=1,
    )
    data = load_data(template(tmp_path, spec))
    assert data.trips == {"6a": {34}, "6b": {35}, "9a": {35}}
    assert data.blackouts == frozenset({36})
    assert data.peer_cap == 3
    assert data.group_week_cap == 1
    assert data.open_weeks == (34, 35)  # only the tutored classes open a room


def test_a_filled_template_plans(tmp_path):
    """A hand-marked template is a solvable workbook."""

    spec = Spec(
        tutored=[("6a", 2), ("6b", 1)],
        peer_classes=["9a"],
        topics=["Reading"],
        weeks=[34, 35, 36],
        trips={"6a": [34, 36], "6b": [35]},
        peer_rows=[("Sander", "Reading", "9a", "", "")],
    )
    data = load_data(template(tmp_path, spec))
    plan = solve(data)
    assert plan.complete
    assert sorted(s.group.name for s in plan.sessions) == ["6a.1", "6a.2", "6b.1"]
    # A group is blocked while its own class travels, so 6a uses the week 6b opens.
    assert {s.group.name: s.week for s in plan.sessions} == {"6a.1": 35, "6a.2": 35, "6b.1": 34}


def test_an_explicit_peers_sheet_is_taken_as_is(tmp_path):
    spec = Spec(
        tutored=[("6a", 1), ("6b", 1)],
        peer_classes=["9a", "9b"],
        topics=["Reading"],
        weeks=[34, 35],
        trips={"6a": [34], "6b": [35]},
        peer_rows=[("Sander; Weise", "Reading", "9a; 9b", "34-35", "2")],
    )
    data = load_data(template(tmp_path, spec))
    assert [(p.name, p.topic, p.classes, p.max_sessions) for p in data.peers] == [
        ("Sander; Weise", "Reading", ("9a", "9b"), 2)
    ]


def test_it_refuses_an_empty_description(tmp_path):
    empty = Spec(tutored=[], peer_classes=[], topics=[], weeks=[])
    with pytest.raises(ValueError):
        write_scaffold(tmp_path / "nope.xlsx", empty)


@pytest.mark.parametrize(
    "items,expected",
    [
        (["6a:2", "6b:3"], [("6a", 2), ("6b", 3)]),
        (["6a"], [("6a", 1)]),
        ([" 6a : 2 "], [("6a", 2)]),
    ],
)
def test_classes_are_parsed(items, expected):
    assert parse_classes(items) == expected


@pytest.mark.parametrize("item", [":2", "6a:0", "6a:x", "6a:2.5"])
def test_bad_class_descriptions_are_rejected(item):
    with pytest.raises(ValueError):
        parse_classes([item])


def test_a_class_listed_twice_is_rejected():
    with pytest.raises(ValueError):
        parse_classes(["6a:2", "6a:1"])


def test_weeks_come_from_dates():
    # Mon 2026-10-26 is ISO week 44; the following Monday is 45.
    start = datetime.date(2026, 10, 26)
    assert week_of(start) == 44
    assert weeks_from_dates(start, 3) == [44, 45, 46]
    # A midweek pick still counts the Monday of its own week.
    assert weeks_from_dates(datetime.date(2026, 10, 30), 1) == [44]
    assert mondays(datetime.date(2026, 10, 30), 2) == [start, datetime.date(2026, 11, 2)]


def test_week_labels_can_be_the_mondays_themselves(tmp_path):
    """A template that runs over a year change is still unambiguous."""

    dates = mondays(datetime.date(2026, 12, 30), 3)  # weeks 53, 1, 2
    spec = Spec(
        tutored=[("6a", 1), ("6b", 1)],
        peer_classes=["9a"],
        topics=["Reading"],
        weeks=[d.isocalendar()[1] for d in dates],
        week_dates=dates,
        trips={"6a": [dates[0].isocalendar()[1]], "6b": [dates[1].isocalendar()[1]]},
    )
    data = load_data(template(tmp_path, spec))
    assert data.weeks == (1, 2, 53)
    assert data.trips["6a"] == {53} and data.trips["6b"] == {1}
    header = [cell.value for cell in load_workbook(template(tmp_path, spec))["Calendar"][1]][1:]
    assert [value.date() for value in header] == dates
