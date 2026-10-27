"""The result sheets: writing a plan twice must not change anything."""

from __future__ import annotations

from openpyxl import load_workbook

from solver import load_workbook as load_data
from solver import solve
from solver.report import RESULT_SHEETS, write_plan


def sheets(path) -> dict[str, list[tuple]]:
    wb = load_workbook(path)
    return {name: list(wb[name].iter_rows(values_only=True)) for name in wb.sheetnames}


def test_result_sheets_are_added(workbook):
    data = load_data(workbook)
    write_plan(workbook, data, solve(data))
    names = load_workbook(workbook).sheetnames
    assert sorted(name for name in names if name in RESULT_SHEETS) == sorted(RESULT_SHEETS)
    assert names[:4] == ["Groups", "Peers", "Calendar", "Settings"]


def test_writing_twice_changes_nothing(workbook):
    """Run the planner over its own output: the sheets must come out identical."""

    data = load_data(workbook)
    write_plan(workbook, data, solve(data))
    first = sheets(workbook)

    again = load_data(workbook)  # the result sheets must not be read as input
    assert [(g.name, g.cls) for g in again.groups] == [(g.name, g.cls) for g in data.groups]
    assert [p.name for p in again.peers] == [p.name for p in data.peers]

    write_plan(workbook, again, solve(again))
    second = sheets(workbook)

    assert first.keys() == second.keys()
    for name, rows in first.items():
        assert rows == second[name], f"sheet '{name}' changed on the second run"


def test_the_input_sheets_survive(workbook):
    before = sheets(workbook)
    data = load_data(workbook)
    write_plan(workbook, data, solve(data))
    after = sheets(workbook)
    for name, rows in before.items():
        assert after[name] == rows


def test_a_copy_can_be_written_elsewhere(workbook, tmp_path):
    data = load_data(workbook)
    target = tmp_path / "plan.xlsx"
    write_plan(workbook, data, solve(data), out=target)
    assert load_workbook(workbook).sheetnames == [
        name for name in load_workbook(target).sheetnames if name not in RESULT_SHEETS
    ]
    assert all(name in load_workbook(target).sheetnames for name in RESULT_SHEETS)


def test_the_plan_grid_shows_every_session(workbook):
    """Two sessions in one week must both be readable, not overwrite each other."""

    data = load_data(workbook)
    plan = solve(data)
    write_plan(workbook, data, plan)
    ws = load_workbook(workbook)["Plan"]
    rows = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    for group in data.groups:
        cells = rows[group.name]
        shown = sum(str(cell).count("/") for cell in cells[2:-2] if cell)
        assert shown == len(data.topics), f"{group.name}: {shown} of {len(data.topics)} shown"
        assert cells[-2] == len(data.topics)
