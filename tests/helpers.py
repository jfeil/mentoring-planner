"""Shared test data and helpers.

The example workbook is small enough to check by hand. Trips, by week 34-40:

    tutored  6a: 34, 38     peers   9a: 37
    tutored  6b: 35, 39     peers  10a: 35, 38
    tutored  6c: 36         peers  11a: 40

A room exists only in weeks where a tutored class travels: 34, 35, 36, 38, 39.
Weeks 37 and 40 only have peer classes away, so they open no room at all.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

FIXTURE = ROOT / "tests" / "fixtures" / "input.xlsx"
EXPECTED = ROOT / "tests" / "fixtures" / "expected.json"

WEEKS = list(range(34, 41))

TRIPS: dict[str, set[int]] = {
    "6a": {34, 38},
    "6b": {35, 39},
    "6c": {36},
    "9a": {37},
    "10a": {35, 38},
    "11a": {40},
}

GROUPS = [("6a.1", "6a"), ("6a.2", "6a"), ("6b.1", "6b"), ("6c.1", "6c")]

PEERS = [
    ("Sander; Weise", "Reading", "9a; 10a", "", ""),
    ("Bergmann", "Reading", "11a", "", ""),
    ("Reuter; Tolles", "Maths", "10a; 9a", "", ""),
    ("Pandur", "Maths", "11a", "", ""),
    ("Simon", "Media", "9a", "", "1"),
    ("Weise", "Media", "10a", "", ""),
]

SETTINGS = [("Max sessions per peer", "4")]


def write_workbook(
    path: Path,
    groups=GROUPS,
    peers=PEERS,
    trips=TRIPS,
    weeks=WEEKS,
    settings=SETTINGS,
    sheet_names=("Groups", "Peers", "Calendar", "Settings"),
    calendar_classes=None,
) -> Path:
    """Write an input workbook; every argument can be replaced by a test."""

    wb = Workbook()
    wb.remove(wb.active)

    if sheet_names[0] is not None:
        ws = wb.create_sheet(sheet_names[0])
        ws.append(["Group", "Class"])
        for row in groups:
            ws.append(list(row))

    if sheet_names[1] is not None:
        ws = wb.create_sheet(sheet_names[1])
        ws.append(["Peers (surnames)", "Topic", "Classes", "Available weeks", "Max sessions"])
        for row in peers:
            ws.append(list(row))

    if sheet_names[2] is not None:
        classes = (
            _classes_in_order(groups, peers) if calendar_classes is None else list(calendar_classes)
        )
        ws = wb.create_sheet(sheet_names[2])
        ws.append(["Class"] + [f"W{w}" for w in weeks])
        for cls in classes:
            ws.append([cls] + ["x" if w in trips.get(cls, set()) else "" for w in weeks])

    if settings is not None and sheet_names[3] is not None:
        ws = wb.create_sheet(sheet_names[3])
        ws.append(["Setting", "Value"])
        for row in settings:
            ws.append(list(row))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _classes_in_order(groups, peers) -> list[str]:
    classes: list[str] = []
    for group in groups:
        if group[1] not in classes:
            classes.append(group[1])
    for peer in peers:
        for cls in peer[2].replace(",", ";").split(";"):
            cls = cls.strip()
            if cls and cls not in classes:
                classes.append(cls)
    return classes


def make_fixture() -> Path:
    """The committed example workbook, rebuilt from the same data."""

    return write_workbook(FIXTURE)


def plan_rows(plan) -> list[dict]:
    """Sessions in a stable, comparable form."""

    return [
        {
            "week": s.week,
            "group": s.group.name,
            "topic": s.peer.topic,
            "peer": s.peer.name,
        }
        for s in sorted(plan.sessions, key=lambda s: (s.group.name, s.peer.topic))
    ]


def expected_document(data, plan) -> dict:
    return {
        "open_weeks": list(data.open_weeks),
        "sessions": plan_rows(plan),
        "peer_loads": {name: count for name, count in sorted(plan.peer_loads().items())},
        "week_loads": {str(w): n for w, n in sorted(plan.week_loads().items())},
        "unresolved": [list(pair) for pair in plan.unresolved],
    }


def write_expected(path: Path, data, plan) -> None:
    path.write_text(json.dumps(expected_document(data, plan), indent=2) + "\n")


def read_expected(path: Path = EXPECTED) -> dict:
    return json.loads(path.read_text())
