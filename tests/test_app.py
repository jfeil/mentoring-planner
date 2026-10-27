"""The browser front end: the same planner, reached over http."""

from __future__ import annotations

import datetime
import io
import threading
import urllib.error
import urllib.parse
import urllib.request
import uuid
from pathlib import Path

import pytest
from openpyxl import load_workbook

from solver import app
from solver import load_workbook as load_data
from solver.scaffold import mondays

START = "2026-10-27"  # a Tuesday in ISO week 44
FORM = {
    "classes": "6a:2 6b:2",
    "peer_classes": "9a 9b",
    "topics": "Reading Maths",
    "start": START,
    "count": "6",
    "trip": ["6a:44", "6b:45", "9a:46"],
    "peer_cap": "4",
    "group_week_cap": "2",
    "blackouts": "48",
}


@pytest.fixture
def base():
    """The server on a free port, shut down after the test."""

    server = app._Server(("127.0.0.1", 0), app._Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_address[1]}"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def get(base: str, path: str) -> tuple[int, str]:
    try:
        with urllib.request.urlopen(base + path) as response:
            return response.status, response.read().decode("utf-8")
    except urllib.error.HTTPError as error:
        return error.code, error.read().decode("utf-8", "replace")


def post(base: str, path: str, fields: dict) -> tuple[int, bytes]:
    body = urllib.parse.urlencode(fields, doseq=True).encode()
    request = urllib.request.Request(
        base + path,
        data=body,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
    )
    try:
        with urllib.request.urlopen(request) as response:
            return response.status, response.read()
    except urllib.error.HTTPError as error:
        return error.code, error.read()


def upload(base: str, path: str, filename: str, data: bytes) -> tuple[int, bytes, str]:
    boundary = uuid.uuid4().hex
    body = io.BytesIO()
    body.write(f"--{boundary}\r\n".encode())
    body.write(
        f'Content-Disposition: form-data; name="workbook"; filename="{filename}"\r\n'
        "Content-Type: application/octet-stream\r\n\r\n".encode()
    )
    body.write(data)
    body.write(f"\r\n--{boundary}--\r\n".encode())
    request = urllib.request.Request(
        base + path,
        data=body.getvalue(),
        headers={"Content-Type": f'multipart/form-data; boundary="{boundary}"'},
    )
    try:
        with urllib.request.urlopen(request) as response:
            return response.status, response.read(), response.headers.get("Content-Disposition") or ""
    except urllib.error.HTTPError as error:
        return error.code, error.read(), ""


def test_the_page_offers_both_jobs(base):
    status, page = get(base, "/")
    assert status == 200
    assert "<title>Mentoring planner</title>" in page
    assert 'action="/plan"' in page and 'action="/template"' in page


def test_an_unknown_address_is_explained(base):
    status, page = get(base, "/nonsense")
    assert status == 404
    assert "Nothing here" in page


def test_the_trip_grid_follows_the_form(base):
    query = urllib.parse.urlencode({k: v for k, v in FORM.items() if k in ("classes", "peer_classes",
                                                                          "topics", "start", "count")})
    status, page = get(base, "/template?" + query)
    assert status == 200
    weeks = [f"W{week}" for week in range(44, 50)]
    assert all(f">{week}<" in page for week in weeks)
    assert "6a.1, 6a.2, 6b.1, 6b.2" in page
    for cls in ("6a", "6b", "9a", "9b"):
        assert f'value="{cls}:' in page
    assert 'name="trip" value="6a:44" checked' in page  # the example marks are pre-ticked


def test_a_bad_class_list_comes_back_readable(base):
    status, page = get(base, "/template?classes=6a:x&topics=Reading&start=" + START + "&count=4")
    assert status == 400
    assert "6a:2" in page


def test_the_generator_writes_a_plannable_workbook(base, tmp_path):
    status, data = post(base, "/template/create", FORM)
    assert status == 200
    with io.BytesIO(data) as handle:
        names = load_workbook(handle).sheetnames
    assert names == ["Groups", "Peers", "Calendar", "Settings", "Read me"]

    path = tmp_path / "generated.xlsx"
    path.write_bytes(data)
    loaded = load_data(path)
    assert [g.name for g in loaded.groups] == ["6a.1", "6a.2", "6b.1", "6b.2"]
    assert loaded.weeks == tuple(range(44, 50))
    assert loaded.trips == {"6a": {44}, "6b": {45}, "9a": {46}, "9b": set()}
    assert loaded.blackouts == frozenset({48})
    assert loaded.peer_cap == 4 and loaded.group_week_cap == 2
    assert loaded.open_weeks == (44, 45)


def test_the_generator_can_plan_in_the_same_step(base):
    status, data = post(base, "/template/create", {**FORM, "plan": "1"})
    assert status == 200
    with io.BytesIO(data) as handle:
        workbook = load_workbook(handle)
    assert "Plan" in workbook.sheetnames
    rows = list(workbook["Sessions"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 8  # four groups, two topics


def test_a_blackout_outside_the_weeks_is_rejected(base):
    status, page = post(base, "/template/create", {**FORM, "blackouts": "12"})
    assert status == 400
    assert "outside the weeks" in page.decode()


def test_a_date_that_is_not_a_date_is_rejected(base):
    status, page = get(base, "/template?classes=6a:1&topics=Reading&start=next+tuesday&count=4")
    assert status == 400
    assert "next tuesday" in page


def test_uploading_a_workbook_returns_it_planned(base, committed_fixture):
    data = Path(committed_fixture).read_bytes()
    status, result, disposition = upload(base, "/plan", "input.xlsx", data)
    assert status == 200
    assert 'filename="input-planned.xlsx"' in disposition
    with io.BytesIO(result) as handle:
        workbook = load_workbook(handle)
    assert "Plan" in workbook.sheetnames
    rows = list(workbook["Sessions"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 12


def test_the_upload_is_read_by_the_same_code_as_the_shell(base, committed_fixture):
    """The browser and the command line must agree, session for session."""

    from solver import solve

    solved = solve(load_data(committed_fixture))
    expected = {(s.week, s.group.name, s.group.cls, s.peer.topic) for s in solved.sessions}
    data = Path(committed_fixture).read_bytes()
    status, result, _ = upload(base, "/plan", "input.xlsx", data)
    with io.BytesIO(result) as handle:
        rows = load_workbook(handle)["Sessions"].iter_rows(min_row=2, values_only=True)
        served = {(row[0], row[1], row[2], row[3]) for row in rows}  # Week, Group, Class, Topic
    assert status == 200
    assert served == expected


def test_a_broken_upload_is_explained(base):
    status, page, _ = upload(base, "/plan", "notes.txt", b"not a workbook")
    assert status == 400
    assert "not an .xlsx workbook" in page.decode()


def test_a_workbook_without_trips_says_what_is_missing(base, tmp_path):
    from helpers import write_workbook

    path = write_workbook(tmp_path / "no-trips.xlsx", trips={})
    status, page, _ = upload(base, "/plan", "no-trips.xlsx", path.read_bytes())
    assert status == 400
    assert "mark the trips in the calendar sheet" in page.decode()


def test_weeks_are_derived_like_the_command_line(base, tmp_path):
    """The date picker and the CLI agree on which ISO weeks a range means."""

    assert mondays(_date(START), 6) == mondays(_date("2026-10-26"), 6)
    status, data = post(base, "/template/create", FORM)
    assert status == 200
    path = tmp_path / "weeks.xlsx"
    path.write_bytes(data)
    assert load_data(path).weeks == tuple(d.isocalendar()[1] for d in mondays(_date(START), 6))


def _date(text: str):
    return datetime.date.fromisoformat(text)
