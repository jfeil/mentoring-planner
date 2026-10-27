"""The command line entry point."""

from __future__ import annotations

import datetime

from main import main
from openpyxl import load_workbook

from helpers import write_workbook
from solver import load_workbook as load_data


def test_it_plans_a_workbook_in_place(tmp_path, capsys):
    path = write_workbook(tmp_path / "input.xlsx")
    assert main([str(path)]) == 0
    wb = load_workbook(path)
    assert "Plan" in wb.sheetnames and "Sessions" in wb.sheetnames
    rows = list(wb["Sessions"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 12
    assert "12 session(s)" in capsys.readouterr().out


def test_it_reports_an_unreadable_workbook(tmp_path, capsys):
    path = tmp_path / "missing.xlsx"
    assert main([str(path)]) == 2
    assert "cannot find" in capsys.readouterr().err


def test_it_reports_a_broken_workbook(tmp_path, capsys):
    path = write_workbook(tmp_path / "broken.xlsx", groups=[])
    assert main([str(path)]) == 2
    assert "groups sheet" in capsys.readouterr().err or "no groups" in capsys.readouterr().err


def test_it_writes_to_another_file_with_out(tmp_path):
    source = write_workbook(tmp_path / "input.xlsx")
    target = tmp_path / "plan.xlsx"
    assert main([str(source), "--out", str(target)]) == 0
    assert "Plan" in load_workbook(target).sheetnames
    assert "Plan" not in load_workbook(source).sheetnames


def test_the_sample_option_creates_an_input_workbook(tmp_path):
    path = tmp_path / "example.xlsx"
    assert main([str(path), "--sample"]) == 0
    names = load_workbook(path).sheetnames
    assert {"Groups", "Peers", "Calendar", "Settings", "Read me"} <= set(names)
    assert "Plan" not in names


def test_a_template_is_written_by_new(tmp_path, capsys):
    path = tmp_path / "template.xlsx"
    code = main([str(path), "--new", "--classes", "6a:2", "6b:1", "--peer-classes", "9a",
                 "--topics", "Reading", "Maths", "--weeks", "36-40"])
    assert code == 0
    assert "template workbook written" in capsys.readouterr().out
    wb = load_workbook(path)
    assert wb.sheetnames == ["Groups", "Peers", "Calendar", "Settings", "Read me"]
    assert [row[0] for row in wb["Groups"].iter_rows(min_row=2, values_only=True)] == [
        "6a.1", "6a.2", "6b.1",
    ]
    assert [cell.value for cell in wb["Calendar"][1]] == ["Class"] + [f"W{w}" for w in range(36, 41)]
    # The template is a real input workbook: the reader accepts it.
    data = load_data(path)
    assert data.topics == ("Reading", "Maths")


def test_new_needs_classes(tmp_path, capsys):
    assert main([str(tmp_path / "x.xlsx"), "--new", "--topics", "Reading"]) == 2
    assert "--classes" in capsys.readouterr().err


def test_new_rejects_a_bad_week_range(tmp_path, capsys):
    assert main([str(tmp_path / "x.xlsx"), "--new", "--classes", "6a:1", "--weeks", "40"]) == 2
    assert "36-45" in capsys.readouterr().err


def test_new_without_weeks_covers_the_coming_ten(tmp_path):
    path = tmp_path / "soon.xlsx"
    assert main([str(path), "--new", "--classes", "6a:1", "6b:1", "--topics", "Reading"]) == 0
    data = load_data(path)
    assert len(data.weeks) == 10
    # The calendar is headed with the Mondays, so the weeks cannot be ambiguous.
    headers = [cell.value for cell in load_workbook(path)["Calendar"][1]][1:]
    assert all(isinstance(value, datetime.datetime) for value in headers)
    assert [value.isocalendar()[1] for value in headers] == list(data.weeks)


def test_it_exits_1_when_a_session_is_missing(tmp_path, capsys):
    # Media has one peer only, and it travels whenever the room is open.
    path = write_workbook(
        tmp_path / "tight.xlsx",
        peers=[("Sander", "Reading", "9a", "", ""), ("Weise", "Media", "10a; 6a", "", "")],
        trips={"6a": {34, 38}, "6b": {35, 39}, "6c": {36}, "9a": set(), "10a": {34, 35, 36, 38, 39}},
    )
    assert main([str(path)]) == 1
    assert "INCOMPLETE" in capsys.readouterr().err
