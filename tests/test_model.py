"""Reading the workbook: the availability rules and the error messages."""

from __future__ import annotations

import pytest

from helpers import write_workbook
from solver import WorkbookError, load_workbook


def peer(data, name: str):
    return next(p for p in data.peers if p.name == name)


def group(data, name: str):
    return next(g for g in data.groups if g.name == name)


def test_a_room_needs_a_tutored_class_on_a_trip(data):
    assert data.open_weeks == (34, 35, 36, 38, 39)


def test_a_peer_class_trip_opens_no_room(data):
    # W37 only 9a is away, W40 only 11a: both are classes of peers, not tutored ones.
    assert not data.is_open(37)
    assert not data.is_open(40)
    # W36: the tutored class 6c travels, so the room is there.
    assert data.is_open(36)


def test_a_group_is_blocked_by_its_own_class_trip(data):
    assert data.allowed_weeks(group(data, "6a.1"), peer(data, "Pandur")) == (35, 36, 39)


def test_a_peer_is_blocked_by_any_of_its_pupils_classes(data):
    assert peer(data, "Sander; Weise").trip_weeks(data.trips) == {35, 37, 38}
    # 6a.1 can use 35, 36, 39 but Sander/Weise are away in 35 (class 10a travels).
    assert data.allowed_weeks(group(data, "6a.1"), peer(data, "Sander; Weise")) == (36, 39)


def test_a_peer_group_is_several_pupils(data):
    sander = peer(data, "Sander; Weise")
    assert sander.members == ("Sander", "Weise")
    assert sander.classes == ("9a", "10a")


def test_session_limits_are_read(data):
    assert data.peer_cap == 4
    assert peer(data, "Simon").max_sessions == 1
    assert data.cap_of(peer(data, "Simon")) == 1
    assert data.cap_of(peer(data, "Weise")) == 4


def test_available_weeks_narrow_a_peer(tmp_path):
    peers = [("Bergmann", "Reading", "11a", "34-35", "")]
    loaded = load_workbook(write_workbook(tmp_path / "narrow.xlsx", peers=peers))
    assert loaded.allowed_weeks(group(loaded, "6a.1"), peer(loaded, "Bergmann")) == (35,)
    assert loaded.allowed_weeks(group(loaded, "6b.1"), peer(loaded, "Bergmann")) == (34,)


def test_blackout_weeks_kill_a_whole_week(tmp_path):
    path = write_workbook(
        tmp_path / "blackout.xlsx",
        settings=[("Max sessions per peer", "4"), ("Blackout weeks", "34, 38")],
    )
    data = load_workbook(path)
    assert data.open_weeks == (35, 36, 39)
    # 6b is away in 35 and 39, so only the blackout-free 36 is left for it.
    assert data.allowed_weeks(group(data, "6b.1"), peer(data, "Pandur")) == (36,)


def test_settings_sheet_is_optional(tmp_path):
    path = write_workbook(tmp_path / "no_settings.xlsx", settings=None)
    data = load_workbook(path)
    assert data.peer_cap is None
    assert data.cap_of(peer(data, "Simon")) == 1  # its own column still counts


def test_german_headers_are_accepted(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    gruppen = wb.create_sheet("Gruppen")
    gruppen.append(["Gruppe", "Klasse"])
    gruppen.append(["6a.1", "6a"])
    gruppen.append(["6b.1", "6b"])
    mentoren = wb.create_sheet("Mentoren")
    mentoren.append(["Mentoren", "Thema", "Klassen"])
    mentoren.append(["Sander; Weise", "Lesen", "9a; 10a"])
    kalender = wb.create_sheet("Kalender")
    kalender.append(["Klasse", "KW 34", "KW 35", "KW 36"])
    kalender.append(["6a", "x", "", ""])
    kalender.append(["6b", "", "x", ""])
    kalender.append(["9a", "", "", "x"])
    kalender.append(["10a", "", "", ""])
    path = tmp_path / "deutsch.xlsx"
    wb.save(path)

    data = load_workbook(path)
    assert data.topics == ("Lesen",)
    assert data.open_weeks == (34, 35)
    # 6a.1 cannot use W34 (its own class travels), the peers cannot use W36.
    assert data.allowed_weeks(group(data, "6a.1"), peer(data, "Sander; Weise")) == (35,)


def test_a_class_that_is_not_in_the_calendar_is_an_error(tmp_path):
    path = write_workbook(
        tmp_path / "unknown.xlsx",
        groups=[("7z.1", "7z")],
        calendar_classes=["6a", "6b", "6c"],
    )
    with pytest.raises(WorkbookError) as error:
        load_workbook(path)
    assert "7z" in str(error.value)


def test_a_missing_sheet_is_an_error(tmp_path):
    path = write_workbook(
        tmp_path / "no_groups.xlsx",
        sheet_names=(None, "Peers", "Calendar", "Settings"),
    )
    with pytest.raises(WorkbookError) as error:
        load_workbook(path)
    assert "groups" in str(error.value)


def test_a_peer_class_missing_from_the_calendar_is_an_error(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Groups").append(["Group", "Class"])
    wb["Groups"].append(["6a.1", "6a"])
    wb["Groups"].append(["6b.1", "6b"])
    peers = wb.create_sheet("Peers")
    peers.append(["Peers", "Topic", "Classes"])
    peers.append(["Sander", "Reading", "12c"])
    calendar = wb.create_sheet("Calendar")
    calendar.append(["Class", "W34", "W35"])
    calendar.append(["6a", "x", ""])
    calendar.append(["6b", "", "x"])
    path = tmp_path / "peer_class.xlsx"
    wb.save(path)

    with pytest.raises(WorkbookError) as error:
        load_workbook(path)
    assert "12c" in str(error.value)


def test_duplicate_rows_are_an_error(tmp_path):
    path = write_workbook(tmp_path / "dup.xlsx", peers=[("Simon", "Media", "9a", "", "1")] * 2)
    with pytest.raises(WorkbookError) as error:
        load_workbook(path)
    assert "twice" in str(error.value)
